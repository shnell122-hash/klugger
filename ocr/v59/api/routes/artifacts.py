import os
import io
import re
import uuid
import zipfile
from flask import Blueprint, request, jsonify, send_file, Response, session
from tools.db import query, execute

artifacts_bp = Blueprint('artifacts', __name__)

MIMES = {
    'md':   ('text/markdown; charset=utf-8',),
    'html': ('text/html; charset=utf-8',),
    'pdf':  ('application/pdf',),
    'docx': ('application/vnd.openxmlformats-officedocument.wordprocessingml.document',),
}


@artifacts_bp.route('/api/artifacts/<case_id>', methods=['GET'])
def list_artifacts_by_case(case_id):
    """Frontend-compatible: case_id as path param."""
    src_filter    = request.args.get('type', 'all')
    artifact_type = request.args.get('artifact_type', '')
    limit         = min(int(request.args.get('limit', 60)), 200)
    results = []

    if src_filter in ('all', 'user'):
        rows = query(
            """SELECT artifact_id, case_id, filename AS name, mime_type,
                      file_size_bytes, checksum_sha256,
                      CHAR_LENGTH(extracted_text) AS text_len,
                      uploaded_at AS created_at,
                      source, NULL AS artifact_type
               FROM user_artifacts
               WHERE case_id=%s AND (
                   COALESCE(source,'') != 'system'
                   OR mime_type LIKE 'audio/%%'
                   OR mime_type LIKE 'video/%%'
               )
               ORDER BY uploaded_at DESC
               LIMIT %s""",
            (case_id, limit), many=True
        )
        results.extend(rows or [])

    if src_filter in ('all', 'system'):
        # Imágenes generadas por IA (guardadas en user_artifacts con source='system')
        rows = query(
            """SELECT artifact_id, case_id, filename AS artifact_name,
                      'image' AS artifact_type, mime_type, file_size_bytes,
                      uploaded_at AS created_at, 'system' AS source
               FROM user_artifacts
               WHERE case_id=%s AND source='system'
               ORDER BY uploaded_at DESC
               LIMIT %s""",
            (case_id, limit), many=True
        )
        results.extend(rows or [])

        atype_clause = "AND artifact_type=%s" if artifact_type else ""
        params = [case_id]
        if artifact_type:
            params.append(artifact_type)
        params.append(limit)
        rows = query(
            f"""SELECT artifact_id, case_id, artifact_name, artifact_type,
                       mime_type, file_size_bytes, share_slug,
                       created_at, 'system' AS source
                FROM system_artifacts
                WHERE case_id=%s {atype_clause}
                ORDER BY created_at DESC
                LIMIT %s""",
            params, many=True
        )
        results.extend(rows or [])

    results.sort(key=lambda x: str(x.get('created_at', '')), reverse=True)
    return jsonify({"artifacts": results})


@artifacts_bp.route('/api/artifacts/file/<artifact_id>', methods=['GET'])
def serve_artifact_file(artifact_id):
    row = query("SELECT file_path, mime_type, filename FROM user_artifacts WHERE artifact_id=%s", (artifact_id,))
    if not row or not os.path.exists(row['file_path']):
        return jsonify({"error": "Archivo no encontrado"}), 404
    return send_file(row['file_path'], mimetype=row['mime_type'],
                     download_name=row['filename'], as_attachment=False)


@artifacts_bp.route('/api/artifacts/<artifact_id>/degrade', methods=['POST'])
def degrade_image(artifact_id):
    row = query(
        "SELECT file_path, mime_type, source FROM user_artifacts WHERE artifact_id=%s",
        (artifact_id,)
    )
    if not row:
        return jsonify({"error": "Artefacto no encontrado"}), 404
    if row.get('source') != 'system' or not (row.get('mime_type') or '').startswith('image/'):
        return jsonify({"error": "Solo aplica a imágenes generadas por IA"}), 400
    file_path = row.get('file_path')
    if not file_path or not os.path.exists(file_path):
        return jsonify({"error": "Archivo no encontrado en disco"}), 404

    body    = request.get_json(force=True) or {}
    blur    = min(max(float(body.get('blur',    1.4)), 0.0), 5.0)
    quality = min(max(int(  body.get('quality', 48)),  10),  95)

    try:
        import io as _io, hashlib as _hs
        from PIL import Image, ImageFilter
        raw = open(file_path, 'rb').read()
        img = Image.open(_io.BytesIO(raw)).convert('RGB')
        if blur > 0:
            img = img.filter(ImageFilter.GaussianBlur(radius=blur))
        buf = _io.BytesIO()
        img.save(buf, format='JPEG', quality=quality, optimize=True)
        processed = buf.getvalue()
        sha256    = _hs.sha256(processed).hexdigest()
        with open(file_path, 'wb') as fh:
            fh.write(processed)
        execute(
            "UPDATE user_artifacts SET file_size_bytes=%s, checksum_sha256=%s WHERE artifact_id=%s",
            (len(processed), sha256, artifact_id)
        )
        return jsonify({"ok": True})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@artifacts_bp.route('/api/artifacts/<artifact_id>/text', methods=['GET'])
def get_artifact_text(artifact_id):
    row = query("SELECT extracted_text FROM user_artifacts WHERE artifact_id=%s", (artifact_id,))
    if not row:
        row = query("SELECT content AS extracted_text FROM system_artifacts WHERE artifact_id=%s", (artifact_id,))
    if not row:
        return jsonify({"error": "Artefacto no encontrado"}), 404
    return jsonify({"text": row.get('extracted_text', '')})


@artifacts_bp.route('/api/artifacts/<artifact_id>', methods=['DELETE'])
def delete_artifact_compat(artifact_id):
    execute("DELETE FROM system_artifacts WHERE artifact_id=%s", (artifact_id,))
    execute("DELETE FROM user_artifacts WHERE artifact_id=%s", (artifact_id,))
    return jsonify({"status": "deleted"})


@artifacts_bp.route('/api/v1/artifacts', methods=['GET'])
def list_artifacts():
    if 'user_id' not in session:
        return jsonify({'error': 'No autenticado', 'authenticated': False}), 401

    case_id       = request.args.get('case_id', '')
    limit         = min(int(request.args.get('limit', 60)), 200)
    src_filter    = request.args.get('type', 'all')      # all | user | system
    artifact_type = request.args.get('artifact_type', '')

    if not case_id:
        return jsonify({"error": "case_id requerido"}), 400

    results = []

    # --- user_artifacts (subidos por usuario) ---
    if src_filter in ('all', 'user'):
        rows = query(
            """SELECT artifact_id, case_id, filename AS name, mime_type,
                      file_size_bytes, checksum_sha256,
                      CHAR_LENGTH(extracted_text) AS text_len,
                      uploaded_at AS created_at,
                      'user' AS source, NULL AS artifact_type
               FROM user_artifacts
               WHERE case_id=%s AND COALESCE(source,'') != 'system'
               ORDER BY uploaded_at DESC
               LIMIT %s""",
            (case_id, limit), many=True
        )
        results.extend(rows or [])

    # --- system_artifacts + imágenes generadas ---
    if src_filter in ('all', 'system'):
        rows = query(
            """SELECT artifact_id, case_id, filename AS artifact_name,
                      'image' AS artifact_type, mime_type, file_size_bytes,
                      uploaded_at AS created_at, 'system' AS source
               FROM user_artifacts
               WHERE case_id=%s AND source='system'
               ORDER BY uploaded_at DESC
               LIMIT %s""",
            (case_id, limit), many=True
        )
        results.extend(rows or [])

        atype_clause = "AND artifact_type=%s" if artifact_type else ""
        params = [case_id]
        if artifact_type:
            params.append(artifact_type)
        params.append(limit)

        rows = query(
            f"""SELECT artifact_id, case_id, artifact_name, artifact_type,
                       mime_type, file_size_bytes, share_slug,
                       created_at, 'system' AS source
                FROM system_artifacts
                WHERE case_id=%s {atype_clause}
                ORDER BY created_at DESC
                LIMIT %s""",
            params, many=True
        )
        results.extend(rows or [])

    # Ordenar mezclados por fecha
    results.sort(key=lambda x: str(x.get('created_at', '')), reverse=True)
    return jsonify({"artifacts": results})


@artifacts_bp.route('/api/v1/artifacts/<artifact_id>', methods=['GET'])
def get_artifact(artifact_id):
    # Buscar en ambas tablas
    row = query(
        "SELECT * FROM system_artifacts WHERE artifact_id=%s", (artifact_id,)
    )
    if not row:
        row = query(
            "SELECT * FROM user_artifacts WHERE artifact_id=%s", (artifact_id,)
        )
    if not row:
        return jsonify({"error": "Artefacto no encontrado"}), 404
    return jsonify(row)


@artifacts_bp.route('/api/v1/artifacts/<artifact_id>', methods=['DELETE'])
def delete_artifact(artifact_id):
    execute("DELETE FROM system_artifacts WHERE artifact_id=%s", (artifact_id,))
    execute("DELETE FROM user_artifacts WHERE artifact_id=%s", (artifact_id,))
    return jsonify({"status": "deleted"})


@artifacts_bp.route('/api/artifacts/<artifact_id>/download', methods=['GET'])
@artifacts_bp.route('/api/v1/artifacts/<artifact_id>/download', methods=['GET'])
def download_artifact(artifact_id):
    """Descarga un artefacto como .md, .html o .docx según el parámetro fmt."""
    fmt = request.args.get('fmt', 'md')

    # Imágenes generadas (guardadas en user_artifacts con source='system')
    img_row = query(
        "SELECT file_path, mime_type, filename FROM user_artifacts WHERE artifact_id=%s AND source='system'",
        (artifact_id,)
    )
    if img_row and img_row.get('file_path') and os.path.exists(img_row['file_path']):
        return send_file(
            img_row['file_path'],
            mimetype=img_row['mime_type'],
            as_attachment=True,
            download_name=img_row['filename']
        )

    row = query(
        "SELECT artifact_name, artifact_type, content FROM system_artifacts WHERE artifact_id=%s",
        (artifact_id,)
    )
    if not row:
        return jsonify({"error": "Artefacto no encontrado"}), 404

    content    = row.get('content') or ''
    name       = row.get('artifact_name', 'documento').replace(' ', '_')
    is_html    = (row.get('artifact_type') == 'html' or
                  (content.lstrip().startswith('<') and '</html>' in content.lower()))
    if not content:
        return jsonify({"error": "Este artefacto no tiene contenido descargable"}), 404

    try:
        data, ext = _content_to_bytes(content, name, fmt, is_html=is_html)
    except ImportError as e:
        pkg = 'markdown python-docx' if 'docx' in str(e) else 'markdown'
        msg = f"Biblioteca no instalada. En el servidor ejecuta:\n/var/www/catalogos/OCR/v59/venv/bin/pip install {pkg}"
        return Response(msg.encode('utf-8'), status=503,
                        mimetype='text/plain',
                        headers={'Content-Disposition': f'attachment; filename="error_{name}.txt"'})
    except Exception as e:
        return Response(str(e).encode('utf-8'), status=500,
                        mimetype='text/plain',
                        headers={'Content-Disposition': f'attachment; filename="error_{name}.txt"'})

    mime = MIMES.get(ext, ('application/octet-stream', f'{{name}}.{ext}'))[0]
    return Response(
        data,
        mimetype=mime,
        headers={'Content-Disposition': f'attachment; filename="{name}.{ext}"'}
    )


_PDF_CSS = """
  @page {margin:2cm}
  body{font-family:Arial,sans-serif;font-size:11pt;line-height:1.6;color:#111}
  h1{font-size:18pt;border-bottom:2px solid #c47a5a;padding-bottom:4pt;color:#1a202c}
  h2{font-size:14pt;color:#2d3748;margin-top:16pt}
  h3{font-size:12pt;color:#4a5568}
  table{border-collapse:collapse;width:100%;margin:10pt 0}
  td,th{border:1px solid #ccc;padding:6pt 8pt;font-size:10pt}
  th{background:#f5f0eb;font-weight:bold}
  tr:nth-child(even){background:#faf6f1}
  code{background:#f0ede8;padding:1pt 4pt;border-radius:3pt;font-size:9.5pt;font-family:monospace}
  pre{background:#f0ede8;padding:10pt;border-radius:5pt;overflow-x:auto}
  ul,ol{margin:6pt 0;padding-left:20pt}
  hr{border:none;border-top:1px solid #ddd;margin:14pt 0}
"""
_HTML_WRAP = (
    '<!DOCTYPE html><html lang="es"><head><meta charset="UTF-8">'
    '<title>{name}</title><style>{css}</style></head><body>{body}</body></html>'
)


def _md_to_html(content):
    import markdown as md_lib
    return md_lib.markdown(content, extensions=['tables', 'fenced_code', 'nl2br'])


def _content_to_bytes(content, name, fmt, is_html=False):
    """Convierte contenido markdown o HTML al formato pedido."""

    if fmt == 'html':
        if is_html:
            return content.encode('utf-8'), 'html'
        body = _md_to_html(content)
        html = _HTML_WRAP.format(
            name=name, css='body{font-family:Arial,sans-serif;max-width:800px;margin:40px auto;padding:20px;line-height:1.6}h1,h2,h3{color:#1a202c}table{border-collapse:collapse;width:100%}td,th{border:1px solid #ccc;padding:8px}@media print{body{margin:0}}',
            body=body
        )
        return html.encode('utf-8'), 'html'

    if fmt == 'pdf':
        from weasyprint import HTML as WH
        body = content if is_html else _md_to_html(content)
        html = _HTML_WRAP.format(name=name, css=_PDF_CSS, body=body)
        return WH(string=html).write_pdf(), 'pdf'

    if fmt == 'docx':
        return _build_docx(content, name, is_html), 'docx'

    # default: markdown
    if is_html:
        # Strip tags for plain MD download of HTML artifact
        stripped = re.sub(r'<[^>]+>', '', content)
        return stripped.encode('utf-8'), 'md'
    return content.encode('utf-8'), 'md'


def _build_docx(content, name, is_html=False):
    """Genera un DOCX desde markdown o HTML con tablas completas."""
    from docx import Document
    from docx.shared import Pt, RGBColor
    doc = Document()
    doc.add_heading(name.replace('_', ' '), 0)

    if is_html:
        # Intentar conversión HTML→DOCX con htmldocx (tablas, estilos)
        try:
            from htmldocx import HtmlToDocx
            HtmlToDocx().add_html_to_document(content, doc)
        except ImportError:
            # Fallback: extraer texto y tablas manualmente con html.parser
            from html.parser import HTMLParser

            class _TableCollector(HTMLParser):
                def __init__(self):
                    super().__init__()
                    self.blocks = []          # lista de ('text'|'table', data)
                    self._in_table = False
                    self._rows = []
                    self._cur_row = []
                    self._cur_cell = []
                    self._buf = []

                def handle_starttag(self, tag, attrs):
                    if tag == 'table':
                        if self._buf:
                            self.blocks.append(('text', ''.join(self._buf).strip()))
                            self._buf = []
                        self._in_table = True; self._rows = []
                    elif tag == 'tr':
                        self._cur_row = []
                    elif tag in ('td', 'th'):
                        self._cur_cell = []

                def handle_endtag(self, tag):
                    if tag == 'table':
                        self._in_table = False
                        self.blocks.append(('table', self._rows))
                        self._rows = []
                    elif tag == 'tr':
                        if self._cur_row:
                            self._rows.append(self._cur_row)
                    elif tag in ('td', 'th'):
                        self._cur_row.append(''.join(self._cur_cell).strip())
                        self._cur_cell = []

                def handle_data(self, data):
                    if self._in_table:
                        self._cur_cell.append(data)
                    else:
                        self._buf.append(data)

                def close(self):
                    super().close()
                    if self._buf:
                        self.blocks.append(('text', ''.join(self._buf).strip()))

            p = _TableCollector()
            p.feed(content)
            p.close()
            for kind, data in p.blocks:
                if kind == 'text':
                    for line in data.split('\n'):
                        if line.strip():
                            doc.add_paragraph(line.strip())
                elif kind == 'table' and data:
                    ncols = max(len(r) for r in data)
                    t = doc.add_table(rows=len(data), cols=ncols)
                    t.style = 'Table Grid'
                    for ri, row in enumerate(data):
                        for ci, cell in enumerate(row):
                            if ci < ncols:
                                t.rows[ri].cells[ci].text = cell
    else:
        # Markdown: parseo línea a línea con soporte correcto de tablas multi-fila
        in_code_block = False
        code_lines = []
        lines = content.split('\n')
        i = 0
        while i < len(lines):
            line = lines[i].rstrip()

            # Code block toggle
            if line.startswith('```'):
                if in_code_block:
                    p = doc.add_paragraph('\n'.join(code_lines))
                    p.style = doc.styles['No Spacing']
                    for run in p.runs:
                        run.font.name = 'Courier New'
                        run.font.size = Pt(9)
                        run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
                    code_lines = []
                    in_code_block = False
                else:
                    in_code_block = True
                i += 1; continue

            if in_code_block:
                code_lines.append(line)
                i += 1; continue

            stripped = line.strip()

            # Headings
            if stripped.startswith('#### '):
                doc.add_heading(stripped[5:], level=4); i += 1; continue
            if stripped.startswith('### '):
                doc.add_heading(stripped[4:], level=3); i += 1; continue
            if stripped.startswith('## '):
                doc.add_heading(stripped[3:], level=2); i += 1; continue
            if stripped.startswith('# '):
                doc.add_heading(stripped[2:], level=1); i += 1; continue

            # Horizontal rule
            if re.match(r'^[-*_]{3,}$', stripped):
                doc.add_paragraph('─' * 50); i += 1; continue

            # Table: recolectar todas las filas consecutivas → una sola tabla
            if stripped.startswith('|') and stripped.endswith('|'):
                table_rows = []
                while i < len(lines):
                    tr = lines[i].strip()
                    if not (tr.startswith('|') and tr.endswith('|')):
                        break
                    # Ignorar filas separadoras (|---|---|)
                    if not re.match(r'^[\s\-:|]+$', tr.replace('|', '')):
                        cells = [c.strip() for c in tr.strip('|').split('|')]
                        table_rows.append(cells)
                    i += 1
                if table_rows:
                    ncols = max(len(r) for r in table_rows)
                    t = doc.add_table(rows=len(table_rows), cols=ncols)
                    t.style = 'Table Grid'
                    for ri, row in enumerate(table_rows):
                        for ci, cell in enumerate(row):
                            if ci < ncols:
                                t.rows[ri].cells[ci].text = cell
                continue

            # Bullet list
            if re.match(r'^[-*+] ', stripped):
                _add_para_with_inline(doc, stripped[2:], style='List Bullet')
                i += 1; continue

            # Numbered list
            m_ol = re.match(r'^(\d+)\. (.+)', stripped)
            if m_ol:
                _add_para_with_inline(doc, m_ol.group(2), style='List Number')
                i += 1; continue

            # Blockquote
            if stripped.startswith('> '):
                p = _add_para_with_inline(doc, stripped[2:])
                p.paragraph_format.left_indent = Pt(20)
                i += 1; continue

            # Empty line
            if not stripped:
                doc.add_paragraph(''); i += 1; continue

            # Normal paragraph with inline formatting
            _add_para_with_inline(doc, stripped)
            i += 1

    b = io.BytesIO()
    doc.save(b)
    return b.getvalue()


def _add_para_with_inline(doc, text, style=None):
    """Añade un párrafo parseando **bold**, *italic*, `code` inline."""
    from docx.shared import Pt, RGBColor
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    # Parse inline tokens: **bold**, *italic*, `code`
    pattern = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)')
    parts = pattern.split(text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*'):
            run = p.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = p.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
        elif part:
            p.add_run(part)
    return p


@artifacts_bp.route('/api/artifacts/<case_id>/zip', methods=['GET'])
def download_case_zip(case_id):
    """Descarga todos los artefactos generados del expediente en un ZIP."""
    buf = io.BytesIO()
    used_names = {}

    def unique_name(name):
        if name not in used_names:
            used_names[name] = 1
            return name
        used_names[name] += 1
        base, ext = os.path.splitext(name)
        return f"{base}_{used_names[name]}{ext}"

    def fmt_for_type(artifact_type, content, is_html):
        """Elige el formato de descarga según el tipo de artefacto."""
        if is_html or artifact_type == 'html':
            return 'html'
        if artifact_type in ('contract', 'brief'):
            return 'docx'
        return 'md'

    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zf:
        # Artefactos de texto generados por Claude
        rows = query(
            """SELECT artifact_name, artifact_type, content
               FROM system_artifacts
               WHERE case_id=%s AND COALESCE(artifact_type,'') != 'session_notes'
               ORDER BY created_at ASC""",
            (case_id,), many=True
        ) or []
        for r in rows:
            safe = (r.get('artifact_name') or 'documento').replace('/', '_').replace('\\', '_')
            content_r = r.get('content') or ''
            is_html_r = (r.get('artifact_type') == 'html' or
                         (content_r.lstrip().startswith('<') and '</html>' in content_r.lower()))
            fmt = fmt_for_type(r.get('artifact_type', ''), content_r, is_html_r)
            try:
                data, ext = _content_to_bytes(content_r, safe, fmt, is_html=is_html_r)
            except Exception:
                data, ext = content_r.encode('utf-8'), 'md'
            fname = unique_name(f"{safe}.{ext}")
            zf.writestr(fname, data)

        # Imágenes y archivos generados por IA (user_artifacts source='system')
        img_rows = query(
            "SELECT file_path, filename FROM user_artifacts WHERE case_id=%s AND source='system' ORDER BY uploaded_at ASC",
            (case_id,), many=True
        ) or []
        for r in img_rows:
            fp = r.get('file_path', '')
            if fp and os.path.exists(fp):
                fname = unique_name(r.get('filename') or os.path.basename(fp))
                zf.write(fp, fname)

    buf.seek(0)
    case_row = query("SELECT case_name FROM cases WHERE case_id=%s", (case_id,))
    case_name = (case_row or {}).get('case_name', case_id)
    zip_name = f"{case_name.replace(' ', '_')}.zip"
    return send_file(buf, mimetype='application/zip', as_attachment=True, download_name=zip_name)


@artifacts_bp.route('/caso/<path:slug>', methods=['GET'])
def serve_shared_html(slug):
    """
    URL pública y persistente para compartir un HTML generado.
    Soporta:  /caso/mi-analisis
              /caso/mi-analisis.html
    No requiere autenticación — accesible desde cualquier navegador.
    """
    # Quitar extensión .html si viene en la URL
    if slug.endswith('.html'):
        slug = slug[:-5]

    row = query(
        """SELECT content, artifact_name, artifact_type, mime_type, case_id
           FROM system_artifacts
           WHERE share_slug = %s""",
        (slug,)
    )
    if not row:
        return (
            "<!DOCTYPE html><html><head><meta charset='UTF-8'>"
            "<title>No encontrado</title></head><body style='font-family:sans-serif;padding:40px'>"
            "<h2>Documento no encontrado</h2>"
            "<p>El enlace puede haber sido eliminado o la URL es incorrecta.</p>"
            "</body></html>",
            404,
            {'Content-Type': 'text/html; charset=utf-8'}
        )

    content    = row.get('content', '')
    art_type   = row.get('artifact_type', '')
    art_name   = row.get('artifact_name', 'Documento')

    # HTML directo — servir tal cual
    if art_type == 'html' or (row.get('mime_type') or '').startswith('text/html'):
        return Response(content, mimetype='text/html; charset=utf-8')

    # Markdown / texto — envoltura HTML mínima para lectura
    html = (
        f"<!DOCTYPE html><html><head>"
        f"<meta charset='UTF-8'><meta name='viewport' content='width=device-width,initial-scale=1'>"
        f"<title>{art_name}</title>"
        f"<style>body{{font-family:Georgia,serif;max-width:860px;margin:40px auto;padding:0 20px;"
        f"color:#1a1a1a;line-height:1.7}}h1,h2,h3{{font-family:sans-serif}}pre{{background:#f4f4f4;"
        f"padding:12px;border-radius:4px;overflow:auto}}code{{background:#f4f4f4;padding:1px 4px}}"
        f"</style></head><body>"
        f"<h1>{art_name}</h1><pre style='white-space:pre-wrap'>{content}</pre>"
        f"</body></html>"
    )
    return Response(html, mimetype='text/html; charset=utf-8')


@artifacts_bp.route('/api/artifacts/search', methods=['POST'])
def search_artifacts_content():
    """Búsqueda full-text de artefactos para agregar al contexto desde el chat."""
    body    = request.get_json(force=True, silent=True) or {}
    case_id = body.get('case_id', '').strip()
    q       = body.get('q', '').strip()
    if not case_id or not q:
        return jsonify({"results": [], "count": 0})
    like = f'%{q}%'
    sys_rows = query(
        """SELECT artifact_id, artifact_name AS name
           FROM system_artifacts
           WHERE case_id=%s AND (artifact_name LIKE %s OR content LIKE %s)
           ORDER BY created_at DESC LIMIT 50""",
        (case_id, like, like), many=True
    ) or []
    usr_rows = query(
        """SELECT artifact_id, filename AS name
           FROM user_artifacts
           WHERE case_id=%s AND (filename LIKE %s OR COALESCE(extracted_text,'') LIKE %s)
           ORDER BY uploaded_at DESC LIMIT 50""",
        (case_id, like, like), many=True
    ) or []
    results = [{'artifact_id': r['artifact_id'], 'name': r['name']}
               for r in sys_rows + usr_rows]
    return jsonify({"results": results, "count": len(results)})


def _make_slug(name: str) -> str:
    s = name.lower().strip()
    for src, dst in [('á','a'),('é','e'),('í','i'),('ó','o'),('ú','u'),
                     ('ñ','n'),('ü','u'),('à','a'),('è','e'),('ì','i'),
                     ('ò','o'),('ù','u')]:
        s = s.replace(src, dst)
    s = re.sub(r'[^a-z0-9]+', '-', s).strip('-')
    return s[:100] or 'artefacto'


def _unique_slug_art(base: str) -> str:
    slug, suffix = base, 2
    while True:
        if not query("SELECT 1 FROM system_artifacts WHERE share_slug=%s", (slug,)):
            return slug
        slug = f"{base}-{suffix}"; suffix += 1
        if suffix > 999:
            return f"{base}-{uuid.uuid4().hex[:6]}"

@artifacts_bp.route('/api/artifacts/<artifact_id>/share-slug', methods=['GET'])
def get_share_slug(artifact_id):
    """Retorna (y genera si falta) el share_slug de un artefacto."""
    row = query(
        "SELECT share_slug, artifact_name FROM system_artifacts WHERE artifact_id=%s",
        (artifact_id,)
    )
    if not row:
        return jsonify({"error": "Artefacto no encontrado"}), 404

    slug = row.get('share_slug')
    if not slug:
        # Generar slug on-the-fly para artefactos creados antes del deploy
        slug = _unique_slug_art(_make_slug(row.get('artifact_name') or 'documento'))
        execute(
            "UPDATE system_artifacts SET share_slug=%s WHERE artifact_id=%s",
            (slug, artifact_id)
        )

    return jsonify({"share_slug": slug, "artifact_name": row.get('artifact_name')})


@artifacts_bp.route('/api/v1/artifacts/save', methods=['POST'])
def save_artifact():
    """Guardar artefacto generado por Claude (tool call save_artifact)."""
    body = request.get_json(force=True, silent=True) or {}
    required = ['case_id', 'artifact_name', 'artifact_type', 'content']
    for field in required:
        if not body.get(field):
            return jsonify({"error": f"{field} requerido"}), 400

    art_id = str(uuid.uuid4())
    content = body['content']
    execute(
        """INSERT INTO system_artifacts
           (artifact_id, case_id, artifact_name, artifact_type, content,
            mime_type, file_size_bytes, source, selected_ctx_json)
           VALUES (%s,%s,%s,%s,%s,'text/markdown',%s,'system',%s)""",
        (
            art_id,
            body['case_id'],
            body['artifact_name'],
            body['artifact_type'],
            content,
            len(content.encode()),
            body.get('selected_ctx_json'),
        )
    )
    return jsonify({"artifact_id": art_id, "status": "saved"}), 201
